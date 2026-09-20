"""Prepara una carga demostrativa, trazable y reversible desde los libros fuente.

El script no modifica los Excel ni se conecta a Cloudflare. Genera dos archivos SQL:
uno de carga y otro de retiro para ejecutar mediante Wrangler cuando corresponda.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import uuid
from collections import defaultdict
from pathlib import Path
from typing import Any

import openpyxl


BATCH_ID = "excel-demo-2026-09-19"
NOW = "2026-09-19T00:00:00.000Z"
FILES = (
    "CONTROL GUIAS.xlsx",
    "CONTROL TRANSPORTE.xlsx",
    "CUADRO DE DESCUENTO DE TRANSPORTES.xlsx",
    "CONTROL FACTURAS.xlsx",
)
PLANT_ALIASES = {
    "ANALYTICA": "ANALYTICA",
    "ANALITYCA": "ANALYTICA",
    "COLIBRI": "COLIBRI",
    "AEQUUM": "AEQUUM",
}
PLANT_NAMES = {
    "ANALYTICA": "Analytica",
    "COLIBRI": "Colibrí",
    "AEQUUM": "Aequum",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def json_value(value: Any) -> Any:
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def json_row(values: list[Any]) -> list[Any]:
    return [json_value(value) for value in values]


def parse_date(value: Any) -> str | None:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    text = clean(value)
    if not text:
        return None
    match = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    return match.group(0) if match else None


def parse_number(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = clean(value).replace(",", "")
    if not text or text in {"-", "–"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def cents(value: Any) -> int | None:
    number = parse_number(value)
    return None if number is None else round(number * 100)


def normalize_gre(value: Any) -> str | None:
    text = clean(value).upper()
    match = re.search(r"\b(EG\d{2})\s*[- ]\s*0*(\d+)\b", text)
    if not match:
        return None
    return f"{match.group(1)}-{int(match.group(2))}"


def normalize_invoice(value: Any) -> str | None:
    text = clean(value).upper()
    if not re.match(r"^[A-Z]\d{3}\s*-\s*\d+", text):
        return None
    return re.sub(r"\s*-\s*", "-", text)


def normalize_lot(value: Any) -> str | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(int(value)) if float(value).is_integer() else str(value)
    text = clean(value).upper()
    if not text or text in {"-", "–"}:
        return None
    return text


def canonical_name(value: Any) -> str | None:
    text = clean(value).upper()
    if not text or text in {"-", "–"}:
        return None
    return text


def sql(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def insert(table: str, values: dict[str, Any]) -> str:
    columns = ", ".join(values)
    rendered = ", ".join(sql(value) for value in values.values())
    return f"INSERT INTO {table} ({columns}) VALUES ({rendered});"


class ImportPlan:
    def __init__(self, batch_id: str, now: str) -> None:
        self.batch_id = batch_id
        self.now = now
        self.statements: list[tuple[str, str]] = []
        self.source_rows: dict[tuple[str, str, int], str] = {}
        self.entity_maps: set[tuple[str, str, str | None]] = set()
        self.issues: list[dict[str, Any]] = []
        self.entity_count = 0

    def uid(self, key: str) -> str:
        return str(uuid.uuid5(uuid.NAMESPACE_URL, f"minantaya:{self.batch_id}:{key}"))

    def add_statement(self, table: str, values: dict[str, Any]) -> None:
        self.statements.append((table, insert(table, values)))

    def map_entity(self, entity_type: str, entity_id: str, source_row_id: str | None = None, source_key: str | None = None) -> None:
        identity = (entity_type, entity_id, source_row_id)
        if identity in self.entity_maps:
            return
        self.entity_maps.add(identity)
        self.add_statement(
            "test_data_entity_map",
            {
                "id": self.uid(f"map:{entity_type}:{entity_id}:{source_row_id or source_key or ''}"),
                "batch_id": self.batch_id,
                "entity_type": entity_type,
                "entity_id": entity_id,
                "source_row_id": source_row_id,
                "source_key": source_key,
                "created_at": self.now,
            },
        )

    def issue(self, severity: str, category: str, description: str, source: tuple[str, str, int] | None = None, details: dict[str, Any] | None = None) -> None:
        source_file, sheet_name, row_number = source if source else (None, None, None)
        position = ":".join(str(piece) for piece in source) if source else category
        self.issues.append(
            {
                "id": self.uid(f"issue:{category}:{position}:{description}"),
                "batch_id": self.batch_id,
                "severity": severity,
                "category": category,
                "source_file": source_file,
                "sheet_name": sheet_name,
                "source_row_number": row_number,
                "description": description,
                "details_json": json.dumps(details or {}, ensure_ascii=False, sort_keys=True),
                "status": "ABIERTA",
                "created_at": self.now,
                "updated_at": self.now,
            },
        )


def read_source_books(source_dir: Path, plan: ImportPlan) -> tuple[dict[str, dict[str, list[tuple[int, list[Any]]]]], list[dict[str, Any]]]:
    books: dict[str, dict[str, list[tuple[int, list[Any]]]]] = {}
    manifest: list[dict[str, Any]] = []
    for file_name in FILES:
        path = source_dir / file_name
        payload = path.read_bytes()
        manifest.append({"file": file_name, "sha256": hashlib.sha256(payload).hexdigest()})
        formulas_book = openpyxl.load_workbook(path, data_only=False, read_only=True)
        values_book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheets: dict[str, list[tuple[int, list[Any]]]] = {}
        for formula_sheet, values_sheet in zip(formulas_book.worksheets, values_book.worksheets, strict=True):
            rows: list[tuple[int, list[Any]]] = []
            formula_rows = list(formula_sheet.iter_rows(values_only=False))
            value_rows = list(values_sheet.iter_rows(values_only=True))
            for row_number, (formula_cells, values) in enumerate(zip(formula_rows, value_rows, strict=True), start=1):
                formula_values = [cell.value for cell in formula_cells]
                if not any(clean(value) for value in values) and not any(clean(value) for value in formula_values):
                    continue
                source_key = (file_name, values_sheet.title, row_number)
                source_id = plan.uid(f"source:{file_name}:{values_sheet.title}:{row_number}")
                value_json = json_row(list(values))
                formulas = {
                    str(index + 1): json_value(value)
                    for index, value in enumerate(formula_values)
                    if isinstance(value, str) and value.startswith("=")
                }
                checksum = hashlib.sha256(
                    json.dumps({"values": value_json, "formulas": formulas}, ensure_ascii=False, sort_keys=True).encode("utf-8"),
                ).hexdigest()
                plan.source_rows[source_key] = source_id
                plan.add_statement(
                    "test_data_source_rows",
                    {
                        "id": source_id,
                        "batch_id": plan.batch_id,
                        "source_file": file_name,
                        "sheet_name": values_sheet.title,
                        "source_row_number": row_number,
                        "values_json": json.dumps(value_json, ensure_ascii=False),
                        "formulas_json": json.dumps(formulas, ensure_ascii=False, sort_keys=True),
                        "row_checksum": checksum,
                        "created_at": plan.now,
                    },
                )
                rows.append((row_number, list(values)))
            sheets[values_sheet.title] = rows
        books[file_name] = sheets
    return books, manifest


def source_id(plan: ImportPlan, file_name: str, sheet_name: str, row_number: int) -> str:
    return plan.source_rows[(file_name, sheet_name, row_number)]


def parse_guide_groups(rows: list[tuple[int, list[Any]]]) -> list[dict[str, Any]]:
    groups: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    current_plant: str | None = None
    current_date: str | None = None
    for row_number, row in rows:
        values = row + [None] * max(0, 13 - len(row))
        if "REGISTRO DE EMISION" in clean(values[1]).upper():
            current = None
            current_plant = None
            current_date = None
            continue
        plant = PLANT_ALIASES.get(clean(values[0]).upper())
        if plant:
            current_plant = plant
        row_date = parse_date(values[1])
        if row_date:
            current_date = row_date
        gre = normalize_gre(values[2])
        if gre:
            current = {
                "gre": gre,
                "original": clean(values[2]),
                "plant": plant or current_plant,
                "issued_at": row_date or current_date,
                "transport": clean(values[3]),
                "rows": [],
            }
            groups.append(current)
        if current is not None:
            current["rows"].append((row_number, values[:13]))
    return groups


def split_suppliers(value: Any) -> list[str]:
    name = canonical_name(value)
    if not name:
        return []
    return [piece.strip() for piece in name.split("/") if piece.strip()]


def guide_status(group: dict[str, Any]) -> tuple[str, str | None]:
    source_text = " ".join(clean(cell).upper() for _, row in group["rows"] for cell in row)
    if "DADA DE BAJA" in source_text:
        return "ANULADA", "La fuente Excel indica DADA DE BAJA."
    if "FACTURADO" in source_text or any(normalize_invoice(row[8]) for _, row in group["rows"]):
        return "FACTURADA", None
    if any(parse_date(row[7]) for _, row in group["rows"]):
        return "LIQUIDADA", None
    return "EMITIDA", None


def extract_invoice_lot_references(group: dict[str, Any], lot_by_code: dict[str, str]) -> dict[str, set[str]]:
    links: dict[str, set[str]] = defaultdict(set)
    active_invoice: str | None = None
    for _, row in group["rows"]:
        if invoice := normalize_invoice(row[8]):
            active_invoice = invoice
        lot_code = normalize_lot(row[5])
        if active_invoice and lot_code and lot_code in lot_by_code:
            links[active_invoice].add(lot_by_code[lot_code])
    return links


def parse_transport_gres(value: Any) -> list[str]:
    text = clean(value).upper()
    result: list[str] = []
    for prefix, start, end in re.findall(r"\b(EG\d{2})\s*[- ]\s*0*(\d+)(?:\s*-\s*0*(\d+))?", text):
        first = int(start)
        last = int(end) if end else first
        if last < first or last - first > 30:
            continue
        result.extend(f"{prefix}-{number}" for number in range(first, last + 1))
    return result


def ensure_counterparty(
    plan: ImportPlan,
    counterparties: dict[tuple[str, str], dict[str, Any]],
    counterparty_type: str,
    name: str,
    source_row_id: str | None,
    document_number: str | None = None,
) -> str:
    key = (counterparty_type, name)
    record = counterparties.get(key)
    if not record:
        record = {
            "id": plan.uid(f"counterparty:{counterparty_type}:{name}"),
            "type": counterparty_type,
            "legal_name": name.title() if name == name.upper() else name,
            "document_number": document_number,
            "created_at": plan.now,
            "updated_at": plan.now,
        }
        counterparties[key] = record
    elif document_number and not record["document_number"]:
        record["document_number"] = document_number
    plan.map_entity("counterparty", record["id"], source_row_id, f"{counterparty_type}:{name}")
    return record["id"]


def build_plan(source_dir: Path, batch_id: str, now: str) -> tuple[ImportPlan, list[dict[str, Any]]]:
    plan = ImportPlan(batch_id, now)
    books, manifest = read_source_books(source_dir, plan)
    counterparties: dict[tuple[str, str], dict[str, Any]] = {}
    plant_ids: dict[str, str] = {}
    for code, legal_name in PLANT_NAMES.items():
        plant_id = plan.uid(f"plant:{code}")
        plant_ids[code] = plant_id
        plan.add_statement("plants", {"id": plant_id, "code": code, "legal_name": legal_name, "active": 1, "created_at": now, "updated_at": now})
        plan.map_entity("plant", plant_id, source_key=code)

    guide_file = "CONTROL GUIAS.xlsx"
    guide_sheet = "📋 Control General"
    guide_groups = parse_guide_groups(books[guide_file][guide_sheet])
    guide_ids: dict[str, str] = {}
    lot_ids: dict[str, str] = {}
    guide_lot_ids: dict[tuple[str, str], str] = {}
    guide_invoice_lots: dict[str, set[str]] = defaultdict(set)
    lot_billed_by: dict[str, str] = {}
    assay_report_ids: set[str] = set()

    for group in guide_groups:
        first_row_number = group["rows"][0][0]
        first_source_id = source_id(plan, guide_file, guide_sheet, first_row_number)
        issued_at = group["issued_at"]
        if not issued_at:
            plan.issue("CRITICA", "GUIA_SIN_FECHA", f"La guía {group['original']} no tiene fecha de emisión utilizable.", (guide_file, guide_sheet, first_row_number))
            continue
        status, void_reason = guide_status(group)
        transport_name = canonical_name(group["transport"])
        carrier_id: str | None = None
        if transport_name and "DADA DE BAJA" not in transport_name and "GUIA PARA" not in transport_name:
            carrier_id = ensure_counterparty(plan, counterparties, "TRANSPORTISTA", transport_name, first_source_id)
        guide_id = plan.uid(f"guide:{group['gre']}")
        guide_ids[group["gre"]] = guide_id
        notes = {"source": {"file": guide_file, "sheet": guide_sheet, "firstRow": first_row_number}}
        plan.add_statement(
            "guides",
            {
                "id": guide_id,
                "gre_original": group["original"],
                "gre_normalized": group["gre"],
                "plant_id": plant_ids.get(group["plant"]),
                "carrier_id": carrier_id,
                "issued_at": f"{issued_at}T00:00:00.000Z",
                "status": status,
                "transport_reference": group["transport"] or None,
                "notes": json.dumps(notes, ensure_ascii=False),
                "voided_at": f"{issued_at}T00:00:00.000Z" if status == "ANULADA" else None,
                "void_reason": void_reason,
                "created_at": now,
                "updated_at": now,
            },
        )
        plan.entity_count += 1
        plan.map_entity("guide", guide_id, first_source_id, group["gre"])
        plan.add_statement(
            "guide_events",
            {
                "id": plan.uid(f"guide-event:{group['gre']}:created"),
                "guide_id": guide_id,
                "event_type": "BAJA" if status == "ANULADA" else "CREADA",
                "occurred_at": f"{issued_at}T00:00:00.000Z",
                "detail_json": json.dumps({"importedFrom": {"file": guide_file, "sheet": guide_sheet, "firstRow": first_row_number}, "status": status}, ensure_ascii=False),
                "created_at": now,
            },
        )
        plan.entity_count += 1
        plan.map_entity("guide_event", plan.uid(f"guide-event:{group['gre']}:created"), first_source_id)

        supplier_names: list[str] = []
        linked_lots: set[str] = set()
        for sequence, (row_number, row) in enumerate(group["rows"], start=1):
            row_source_id = source_id(plan, guide_file, guide_sheet, row_number)
            if row[4] not in (None, "", "-", "–"):
                supplier_names = split_suppliers(row[4])
            lot_code = normalize_lot(row[5])
            if not lot_code:
                continue
            lot_id = lot_ids.get(lot_code)
            if not lot_id:
                lot_id = plan.uid(f"lot:{lot_code}")
                lot_ids[lot_code] = lot_id
                withdrawal = "RETIRO" in clean(row[10]).upper()
                sacks = parse_number(row[6])
                plan.add_statement(
                    "lots",
                    {
                        "id": lot_id,
                        "code": lot_code,
                        "sack_count": int(sacks) if sacks is not None and sacks >= 0 and sacks.is_integer() else None,
                        "status": "RETIRO_PENDIENTE" if withdrawal else "ACTIVO",
                        "created_at": now,
                        "updated_at": now,
                    },
                )
                plan.entity_count += 1
                plan.map_entity("lot", lot_id, row_source_id, lot_code)
            if lot_id in linked_lots:
                plan.issue("INFO", "LOTE_REFERENCIADO_MAS_DE_UNA_VEZ", f"El lote {lot_code} aparece más de una vez en la guía {group['gre']}; se conserva cada fila fuente y se usa una sola relación operativa.", (guide_file, guide_sheet, row_number))
                continue
            linked_lots.add(lot_id)
            guide_lot_id = plan.uid(f"guide-lot:{guide_id}:{lot_id}")
            guide_lot_ids[(group["gre"], lot_code)] = guide_lot_id
            withdrawal = "RETIRO" in clean(row[10]).upper()
            plan.add_statement(
                "guide_lots",
                {
                    "id": guide_lot_id,
                    "guide_id": guide_id,
                    "lot_id": lot_id,
                    "sequence": sequence,
                    "withdrawal_requested": 1 if withdrawal else 0,
                    "withdrawal_reason": clean(row[10]) or None,
                    "created_at": now,
                    "updated_at": now,
                },
            )
            plan.entity_count += 1
            plan.map_entity("guide_lot", guide_lot_id, row_source_id, f"{group['gre']}:{lot_code}")
            for supplier_name in supplier_names:
                supplier_id = ensure_counterparty(plan, counterparties, "PROVEEDOR", supplier_name, row_source_id)
                supplier_link_id = plan.uid(f"lot-supplier:{guide_lot_id}:{supplier_id}")
                plan.add_statement(
                    "lot_suppliers",
                    {"id": supplier_link_id, "guide_lot_id": guide_lot_id, "supplier_id": supplier_id, "created_at": now, "updated_at": now},
                )
                plan.entity_count += 1
                plan.map_entity("lot_supplier", supplier_link_id, row_source_id)
        for invoice_number, lot_set in extract_invoice_lot_references(group, lot_ids).items():
            guide_invoice_lots[invoice_number].update(lot_set)

    commercial_file = "CONTROL FACTURAS.xlsx"
    commercial_invoices: dict[str, dict[str, Any]] = {}
    sheet_priority = {"JULIO": 1, "AGOSTO": 2, "SETIEMBRE": 3}
    for sheet_name in ("JULIO", "AGOSTO", "SETIEMBRE"):
        for row_number, row in books[commercial_file][sheet_name]:
            if row_number == 1:
                continue
            values = row + [None] * max(0, 16 - len(row))
            invoice_number = normalize_invoice(values[2])
            if not invoice_number:
                continue
            row_source_id = source_id(plan, commercial_file, sheet_name, row_number)
            issued_date = parse_date(values[3])
            source_status = clean(values[6] if sheet_name != "JULIO" else values[8]).upper()
            provider = canonical_name(values[1])
            is_void = provider == "ANULADA" or source_status == "ANULADA"
            amount = cents(values[4])
            if amount is None:
                amount = 0 if is_void else None
            if not issued_date or amount is None:
                plan.issue("ADVERTENCIA", "FACTURA_COMERCIAL_INCOMPLETA", f"La factura comercial {invoice_number} no tiene fecha o monto utilizable para una tabla operativa.", (commercial_file, sheet_name, row_number))
                continue
            record = {
                "id": plan.uid(f"commercial-invoice:{invoice_number}"),
                "invoice_number": invoice_number,
                "issued_at": f"{issued_date}T00:00:00.000Z",
                "amount_usd_cents": amount,
                "detraction_percent": 0.10,
                "detraction_pen_cents": 0,
                "status": "ANULADA" if is_void else "PAGADA" if source_status == "PAGADO" else "EMITIDA",
                "void_reason": "La fuente Excel indica ANULADA." if is_void else None,
                "created_at": now,
                "updated_at": now,
            }
            existing = commercial_invoices.get(invoice_number)
            if existing and sheet_priority[sheet_name] < existing["_source_priority"]:
                continue
            record["_source_priority"] = sheet_priority[sheet_name]
            commercial_invoices[invoice_number] = record
            plan.map_entity("commercial_invoice", record["id"], row_source_id, invoice_number)

    for invoice_number, record in commercial_invoices.items():
        plan.add_statement("commercial_invoices", {key: value for key, value in record.items() if key != "_source_priority"})
        plan.entity_count += 1
        linked_lot_ids = set(guide_invoice_lots.get(invoice_number, set()))
        if not linked_lot_ids:
            continue
        available_lots = [lot_id for lot_id in sorted(linked_lot_ids) if lot_id not in lot_billed_by]
        if len(available_lots) != len(linked_lot_ids):
            plan.issue("ADVERTENCIA", "LOTE_FACTURADO_EN_MAS_DE_UNA_FACTURA", f"No se forzó una segunda asignación de lotes para la factura {invoice_number}; la fila fuente permanece íntegra.")
        if not available_lots:
            continue
        base, remainder = divmod(record["amount_usd_cents"], len(available_lots))
        for index, lot_id in enumerate(available_lots):
            link_id = plan.uid(f"commercial-invoice-lot:{record['id']}:{lot_id}")
            plan.add_statement(
                "commercial_invoice_lots",
                {"id": link_id, "commercial_invoice_id": record["id"], "lot_id": lot_id, "amount_usd_cents": base + (1 if index < remainder else 0), "created_at": now},
            )
            plan.entity_count += 1
            plan.map_entity("commercial_invoice_lot", link_id, source_key=invoice_number)
            lot_billed_by[lot_id] = invoice_number

    transport_file = "CONTROL TRANSPORTE.xlsx"
    transport_rows: list[dict[str, Any]] = []
    for sheet_name in ("Control de Pagos", "LOTES VOLADOS "):
        for row_number, row in books[transport_file][sheet_name]:
            if row_number == 1:
                continue
            values = row + [None] * max(0, 17 - len(row))
            invoice_number = normalize_invoice(values[4])
            carrier_name = canonical_name(values[3])
            if not invoice_number or not carrier_name:
                continue
            transport_rows.append({"file": transport_file, "sheet": sheet_name, "row": row_number, "values": values, "invoice": invoice_number, "carrier": carrier_name})

    rates_by_date: dict[str, set[float]] = defaultdict(set)
    rate_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in transport_rows:
        issued_date = parse_date(item["values"][1])
        rate = parse_number(item["values"][8])
        if issued_date and rate and rate > 0:
            rates_by_date[issued_date].add(rate)
            rate_rows[issued_date].append(item)
    exchange_rate_ids: dict[str, str] = {}
    for rate_date, rates in rates_by_date.items():
        if len(rates) > 1:
            for item in rate_rows[rate_date]:
                plan.issue("ADVERTENCIA", "TIPO_CAMBIO_CONFLICTIVO", f"Hay más de un tipo de cambio para {rate_date}; se conserva la fuente sin escoger uno automáticamente.", (item["file"], item["sheet"], item["row"]), {"rates": sorted(rates)})
            continue
        rate = next(iter(rates))
        exchange_id = plan.uid(f"exchange-rate:{rate_date}")
        exchange_rate_ids[rate_date] = exchange_id
        plan.add_statement("exchange_rates", {"id": exchange_id, "rate_date": rate_date, "source": "Importado desde CONTROL TRANSPORTE.xlsx", "usd_to_pen": rate, "created_at": now, "updated_at": now})
        plan.entity_count += 1
        for item in rate_rows[rate_date]:
            plan.map_entity("exchange_rate", exchange_id, source_id(plan, item["file"], item["sheet"], item["row"]), rate_date)

    transport_invoices: dict[tuple[str, str], dict[str, Any]] = {}
    for item in transport_rows:
        values = item["values"]
        source_key = (item["file"], item["sheet"], item["row"])
        row_source_id = source_id(plan, *source_key)
        issued_date = parse_date(values[1])
        amount = cents(values[7])
        if not issued_date or amount is None:
            plan.issue("ADVERTENCIA", "FACTURA_TRANSPORTE_INCOMPLETA", f"La factura de transporte {item['invoice']} no tiene fecha o monto utilizable para una tabla operativa.", source_key)
            continue
        ruc = clean(values[2])
        if ruc.endswith(".0"):
            ruc = ruc[:-2]
        carrier_id = ensure_counterparty(plan, counterparties, "TRANSPORTISTA", item["carrier"], row_source_id, ruc or None)
        payment_status = " ".join(clean(values[index]).upper() for index in (12, 13, 14))
        is_void = "ANUL" in payment_status
        invoice_status = "ANULADA" if is_void else "PAGADA" if "PAGADO" in payment_status else "REGISTRADA"
        key = (carrier_id, item["invoice"])
        record = {
            "id": plan.uid(f"transport-invoice:{carrier_id}:{item['invoice']}"),
            "carrier_id": carrier_id,
            "invoice_number": item["invoice"],
            "issued_at": f"{issued_date}T00:00:00.000Z",
            "amount_usd_cents": amount,
            "exchange_rate_id": exchange_rate_ids.get(issued_date),
            "detraction_percent": parse_number(values[9]) or 0.04,
            "detraction_pen_cents": cents(values[10]) or 0,
            "status": invoice_status,
            "void_reason": "La fuente Excel indica ANULADA." if is_void else None,
            "created_at": now,
            "updated_at": now,
        }
        existing = transport_invoices.get(key)
        if existing and any(existing[field] != record[field] for field in ("issued_at", "amount_usd_cents", "status")):
            plan.issue("CRITICA", "FACTURA_TRANSPORTE_DUPLICADA", f"La factura de transporte {item['invoice']} aparece con datos incompatibles para el mismo transportista.", source_key)
            continue
        transport_invoices[key] = record
        plan.map_entity("transport_invoice", record["id"], row_source_id, f"{item['carrier']}:{item['invoice']}")
        for gre in parse_transport_gres(values[5]):
            guide_id = guide_ids.get(gre)
            if not guide_id:
                plan.issue("ADVERTENCIA", "GUIA_TRANSPORTE_NO_ENCONTRADA", f"La factura de transporte {item['invoice']} referencia la guía {gre}, que no se encontró en CONTROL GUIAS.", source_key)
                continue
            link_id = plan.uid(f"transport-invoice-guide:{record['id']}:{guide_id}")
            plan.add_statement("transport_invoice_guides", {"id": link_id, "transport_invoice_id": record["id"], "guide_id": guide_id, "created_at": now})
            plan.entity_count += 1
            plan.map_entity("transport_invoice_guide", link_id, row_source_id, gre)
        net_amount = cents(values[11])
        if invoice_status == "PAGADA" and net_amount and net_amount > 0:
            payment_id = plan.uid(f"transport-payment:{record['id']}")
            plan.add_statement("payments", {"id": payment_id, "payment_type": "TRANSPORTE", "transport_invoice_id": record["id"], "paid_at": f"{issued_date}T00:00:00.000Z", "currency": "USD", "amount_cents": net_amount, "reference": "Importado como pago neto", "status": "CONFIRMADO", "created_at": now, "updated_at": now})
            plan.entity_count += 1
            plan.map_entity("payment", payment_id, row_source_id, item["invoice"])
        detraction_amount = cents(values[10])
        if clean(values[13]).upper() == "PAGADO" and detraction_amount and detraction_amount > 0:
            payment_id = plan.uid(f"transport-detraction:{record['id']}")
            plan.add_statement("payments", {"id": payment_id, "payment_type": "DETRACCION_TRANSPORTE", "transport_invoice_id": record["id"], "paid_at": f"{issued_date}T00:00:00.000Z", "currency": "PEN", "amount_cents": detraction_amount, "reference": "Importado como detracción", "status": "CONFIRMADO", "created_at": now, "updated_at": now})
            plan.entity_count += 1
            plan.map_entity("payment", payment_id, row_source_id, item["invoice"])

    for record in transport_invoices.values():
        plan.add_statement("transport_invoices", record)
        plan.entity_count += 1

    for counterparty in counterparties.values():
        plan.add_statement(
            "counterparties",
            {**counterparty, "active": 1},
        )
        plan.entity_count += 1

    laboratory_file = "CUADRO DE DESCUENTO DE TRANSPORTES.xlsx"
    laboratory_sheet = "CUADRO LABORATORIO "
    current_gre: str | None = None
    for row_number, row in books[laboratory_file][laboratory_sheet]:
        if row_number <= 3:
            continue
        values = row + [None] * max(0, 6 - len(row))
        if gre := normalize_gre(values[0]):
            current_gre = gre
        lot_code = normalize_lot(values[1])
        if not current_gre or not lot_code:
            continue
        row_source_id = source_id(plan, laboratory_file, laboratory_sheet, row_number)
        guide_id = guide_ids.get(current_gre)
        lot_id = lot_ids.get(lot_code)
        if not guide_id or not lot_id:
            plan.issue("ADVERTENCIA", "RESULTADO_LEYES_SIN_VINCULO", f"El resultado de leyes para {lot_code} no se pudo vincular de forma segura a guía y lote operativos.", (laboratory_file, laboratory_sheet, row_number), {"gre": current_gre})
            continue
        for source_kind, result_index in (("PLANTA", 3), ("EXTERNO", 4)):
            result_value = parse_number(values[result_index])
            if result_value is None:
                continue
            report_id = plan.uid(f"assay-report:{current_gre}:{source_kind}")
            if report_id not in assay_report_ids:
                assay_report_ids.add(report_id)
                plan.add_statement("assay_reports", {"id": report_id, "guide_id": guide_id, "report_number": f"IMPORTADO-{current_gre}-{source_kind}", "reported_at": now, "received_at": now, "status": "RECIBIDO", "source": source_kind, "created_at": now, "updated_at": now})
                plan.entity_count += 1
            result_id = plan.uid(f"assay-result:{report_id}:{lot_id}:LEY")
            plan.add_statement("assay_results", {"id": result_id, "assay_report_id": report_id, "lot_id": lot_id, "element": "LEY", "result_value": result_value, "unit": "LEY", "created_at": now})
            plan.entity_count += 1
            plan.map_entity("assay_report", report_id, row_source_id, f"{current_gre}:{source_kind}")
            plan.map_entity("assay_result", result_id, row_source_id, lot_code)

    for issue in plan.issues:
        plan.add_statement("test_data_quality_issues", issue)
    return plan, manifest


def render_import_sql(plan: ImportPlan, manifest: list[dict[str, Any]]) -> str:
    batch = {
        "id": plan.batch_id,
        "label": "Datos de prueba importados de los libros operativos",
        "is_test_data": 1,
        "source_manifest_json": json.dumps(manifest, ensure_ascii=False, sort_keys=True),
        "source_row_count": len(plan.source_rows),
        "structured_entity_count": plan.entity_count,
        "created_by": "administrador",
        "created_at": plan.now,
    }
    audit = {
        "id": plan.uid("audit:import"),
        "actor_username": "administrador",
        "actor_source": "system",
        "action": "IMPORTED_TEST_DATA",
        "entity_type": "test_data_import_batch",
        "entity_id": plan.batch_id,
        "after_json": json.dumps({"sourceRows": len(plan.source_rows), "structuredEntities": plan.entity_count, "qualityIssues": len(plan.issues), "files": [entry["file"] for entry in manifest]}, ensure_ascii=False),
        "reason": "Carga inicial de datos de prueba desde los Excel entregados por el cliente.",
        "created_at": plan.now,
    }
    table_order = {
        "test_data_source_rows": 1,
        "plants": 2,
        "counterparties": 3,
        "exchange_rates": 4,
        "lots": 5,
        "guides": 6,
        "guide_events": 7,
        "guide_lots": 8,
        "lot_suppliers": 9,
        "assay_reports": 10,
        "assay_results": 11,
        "commercial_invoices": 12,
        "commercial_invoice_lots": 13,
        "transport_invoices": 14,
        "transport_invoice_guides": 15,
        "payments": 16,
        "test_data_entity_map": 17,
        "test_data_quality_issues": 18,
    }
    ordered = [statement for _, (_, statement) in sorted(enumerate(plan.statements), key=lambda item: (table_order[item[1][0]], item[0]))]
    return "\n".join(["PRAGMA foreign_keys = ON;", "BEGIN IMMEDIATE;", insert("test_data_import_batches", batch), *ordered, insert("audit_logs", audit), "COMMIT;", ""])


def render_cleanup_sql(batch_id: str, now: str) -> str:
    def ids(entity_type: str) -> str:
        return f"(SELECT entity_id FROM test_data_entity_map WHERE batch_id = {sql(batch_id)} AND entity_type = {sql(entity_type)})"

    audit = insert(
        "audit_logs",
        {
            "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"minantaya:{batch_id}:audit:cleanup")),
            "actor_username": "administrador",
            "actor_source": "system",
            "action": "REMOVED_TEST_DATA",
            "entity_type": "test_data_import_batch",
            "entity_id": batch_id,
            "after_json": json.dumps({"batchId": batch_id, "preserved": ["users", "auth_sessions", "audit_logs", "plants", "counterparties"]}),
            "reason": "Retiro solicitado de datos de prueba importados desde Excel.",
            "created_at": now,
        },
    )
    deletes = [
        f"DELETE FROM payments WHERE id IN {ids('payment')};",
        f"DELETE FROM documents WHERE id IN {ids('document')};",
        f"DELETE FROM alerts WHERE id IN {ids('alert')};",
        f"DELETE FROM settlement_lines WHERE id IN {ids('settlement_line')};",
        f"DELETE FROM discounts WHERE id IN {ids('discount')};",
        f"DELETE FROM settlements WHERE id IN {ids('settlement')};",
        f"DELETE FROM commercial_invoice_lots WHERE id IN {ids('commercial_invoice_lot')};",
        f"DELETE FROM transport_invoice_guides WHERE id IN {ids('transport_invoice_guide')};",
        f"DELETE FROM assay_results WHERE id IN {ids('assay_result')};",
        f"DELETE FROM assay_reports WHERE id IN {ids('assay_report')};",
        f"DELETE FROM lot_suppliers WHERE id IN {ids('lot_supplier')};",
        f"DELETE FROM guide_lots WHERE id IN {ids('guide_lot')};",
        f"DELETE FROM guide_events WHERE id IN {ids('guide_event')};",
        f"DELETE FROM resamples WHERE id IN {ids('resample')};",
        f"DELETE FROM disputes WHERE id IN {ids('dispute')};",
        f"DELETE FROM purchase_proposals WHERE id IN {ids('purchase_proposal')};",
        f"DELETE FROM commercial_invoices WHERE id IN {ids('commercial_invoice')};",
        f"DELETE FROM transport_invoices WHERE id IN {ids('transport_invoice')};",
        f"DELETE FROM exchange_rates WHERE id IN {ids('exchange_rate')};",
        f"DELETE FROM lots WHERE id IN {ids('lot')};",
        f"DELETE FROM guides WHERE id IN {ids('guide')};",
        f"DELETE FROM test_data_quality_issues WHERE batch_id = {sql(batch_id)};",
        f"DELETE FROM test_data_entity_map WHERE batch_id = {sql(batch_id)};",
        f"DELETE FROM test_data_source_rows WHERE batch_id = {sql(batch_id)};",
        f"DELETE FROM test_data_import_batches WHERE id = {sql(batch_id)};",
    ]
    return "\n".join(["PRAGMA foreign_keys = ON;", "BEGIN IMMEDIATE;", audit, *deletes, "COMMIT;", ""])


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera SQL reversible para los Excel de prueba.")
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--cleanup-output", type=Path, required=True)
    parser.add_argument("--batch-id", default=BATCH_ID)
    parser.add_argument("--now", default=NOW)
    args = parser.parse_args()
    plan, manifest = build_plan(args.source_dir, args.batch_id, args.now)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.cleanup_output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(render_import_sql(plan, manifest), encoding="utf-8")
    args.cleanup_output.write_text(render_cleanup_sql(args.batch_id, args.now), encoding="utf-8")
    print(json.dumps({"batchId": args.batch_id, "sourceRows": len(plan.source_rows), "structuredEntities": plan.entity_count, "qualityIssues": len(plan.issues), "files": manifest}, ensure_ascii=False))


if __name__ == "__main__":
    main()
